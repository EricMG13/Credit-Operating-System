#!/usr/bin/env python3
"""Render a contract-valid CP Markdown handoff into a deterministic DOCX.

The Markdown handoff is authoritative.  The renderer performs no calculation,
rewriting, summarisation, or reconciliation; it only maps Markdown structure to
Word structure.  This makes numerical and semantic parity testable.
"""

from __future__ import annotations

import argparse
import hashlib
import re
import tempfile
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Sequence
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.shared import Inches, Pt, RGBColor, Twips
from docx.table import Table
from docx.text.paragraph import Paragraph

try:
    from tools.validate_handoff import _parse_frontmatter, validate_text
except ModuleNotFoundError:  # Direct execution from tools/.
    from validate_handoff import _parse_frontmatter, validate_text


PORTRAIT_CONTENT_WIDTH_DXA = 9360
LANDSCAPE_CONTENT_WIDTH_DXA = 12960
# Keep the fixed table width exactly equal to the section's writable width.
# A non-zero indent makes the effective footprint wider than the page grid and
# causes LibreOffice to displace cells in some overflow-table projections.
TABLE_INDENT_DXA = 0
CELL_MARGINS_DXA = {"top": 80, "bottom": 80, "start": 120, "end": 120}
FIXED_TIMESTAMP = datetime(2026, 7, 21, 0, 0, 0, tzinfo=timezone.utc)
WIDE_TABLE_COLUMN_LIMIT = 6
LANDSCAPE_TABLE_COLUMN_LIMIT = 12


@dataclass(frozen=True)
class Block:
    kind: str
    value: Any


def clean_inline(text: str) -> str:
    """Return display text without changing analytical tokens."""

    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    text = text.replace("**", "").replace("__", "")
    text = text.replace("`", "")
    return text.strip()


def _split_table_row(line: str) -> list[str]:
    return [clean_inline(cell.strip()) for cell in line.strip().strip("|").split("|")]


def _is_table_separator(line: str) -> bool:
    cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def parse_markdown_body(body: str) -> list[Block]:
    """Parse the deliberately small Markdown subset used by CP handoffs."""

    lines = body.splitlines()
    blocks: list[Block] = []
    paragraph: list[str] = []

    def flush_paragraph() -> None:
        if paragraph:
            blocks.append(Block("paragraph", clean_inline(" ".join(paragraph))))
            paragraph.clear()

    index = 0
    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()
        if not stripped:
            flush_paragraph()
            index += 1
            continue

        if re.fullmatch(
            r"<!--\s*table-id:\s*[A-Za-z][A-Za-z0-9._:-]{2,80}\s*-->",
            stripped,
        ):
            flush_paragraph()
            index += 1
            continue

        heading = re.fullmatch(r"(#{2,4})\s+(.+?)\s*", stripped)
        if heading:
            flush_paragraph()
            blocks.append(Block(f"h{len(heading.group(1))}", clean_inline(heading.group(2))))
            index += 1
            continue

        if (
            stripped.startswith("|")
            and index + 1 < len(lines)
            and _is_table_separator(lines[index + 1])
        ):
            flush_paragraph()
            rows = [_split_table_row(stripped)]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(_split_table_row(lines[index]))
                index += 1
            width = len(rows[0])
            if width == 0 or any(len(row) != width for row in rows):
                raise ValueError("Markdown table rows have inconsistent widths")
            blocks.append(Block("table", rows))
            continue

        bullet = re.fullmatch(r"[-*]\s+(.+)", stripped)
        if bullet:
            flush_paragraph()
            blocks.append(Block("bullet", clean_inline(bullet.group(1))))
            index += 1
            continue

        numbered = re.fullmatch(r"\d+[.)]\s+(.+)", stripped)
        if numbered:
            flush_paragraph()
            blocks.append(Block("number", clean_inline(numbered.group(1))))
            index += 1
            continue

        paragraph.append(stripped)
        index += 1

    flush_paragraph()
    return blocks


def serialise_blocks(blocks: Sequence[Block]) -> list[str]:
    """Create the ordered semantic token stream used for parity testing."""

    serialised: list[str] = []
    for block in blocks:
        if block.kind.startswith("h"):
            serialised.append(f"{block.kind.upper()}:{block.value}")
        elif block.kind == "table":
            serialised.extend("T:" + "\t".join(row) for row in block.value)
        elif block.kind in {"bullet", "number", "paragraph"}:
            serialised.append(f"P:{block.value}")
        else:
            raise ValueError(f"Unsupported block kind: {block.kind}")
    return serialised


def _set_run_font(run, *, name: str = "Calibri", size: float | None = None,
                  bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def _set_cell_margins(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, width in CELL_MARGINS_DXA.items():
        element = tc_mar.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(width))
        element.set(qn("w:type"), "dxa")


def _column_widths(rows: Sequence[Sequence[str]], content_width_dxa: int) -> list[int]:
    maxima = [max(4, max(len(row[index]) for row in rows)) for index in range(len(rows[0]))]
    weights: list[int] = []
    for index, maximum in enumerate(maxima):
        header = rows[0][index].strip().lower()
        minimum = 6
        if header in {"value", "fy2025 value", "amount", "score"}:
            minimum = 18
        elif header in {"unit", "result", "status", "lineage"}:
            minimum = 14
        elif header.endswith("id") or header == "trap":
            minimum = 10
        weights.append(min(36, max(minimum, maximum)))
    total = sum(weights)
    widths = [round(content_width_dxa * weight / total) for weight in weights]
    widths[-1] += content_width_dxa - sum(widths)
    return widths


def _apply_table_geometry(table, widths: Sequence[int], content_width_dxa: int) -> None:
    if sum(widths) != content_width_dxa:
        raise ValueError(f"Column widths must sum to {content_width_dxa} DXA")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    def width_element(tag: str, value: int):
        element = tbl_pr.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tbl_pr.append(element)
        element.set(qn("w:type"), "dxa")
        element.set(qn("w:w"), str(value))

    width_element("w:tblW", content_width_dxa)
    width_element("w:tblInd", TABLE_INDENT_DXA)
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        row.height = None
        for index, cell in enumerate(row.cells):
            cell.width = Twips(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(widths[index]))
            _set_cell_margins(cell)


def _repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def _prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def _configure_styles(document: DocumentObject) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    settings = {
        "Heading 1": (16, "2E74B5", 16, 8),
        "Heading 2": (13, "2E74B5", 12, 6),
        "Heading 3": (12, "1F4D78", 8, 4),
    }
    for name, (size, color, before, after) in settings.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def _add_title_block(document: DocumentObject, fields: dict[str, Any], source_hash: str) -> None:
    kicker = document.add_paragraph()
    kicker.paragraph_format.space_after = Pt(5)
    _set_run_font(kicker.add_run("CP ANALYTICAL REPORT"), size=9, bold=True, color="2E74B5")

    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    _set_run_font(
        title.add_run(f"{fields['module_id']} {fields['module_name']}"),
        size=23,
        bold=True,
        color="000000",
    )

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    subject_name = fields.get("subject_name", fields.get("issuer_name", "[Unknown subject]"))
    _set_run_font(subtitle.add_run(str(subject_name)), size=14, color="373737")

    metadata = (
        ("Reporting period", fields["reporting_period"]),
        ("Analysis date", fields["analysis_date"]),
        ("Run ID", fields["run_id"]),
        ("Confidence", f"{fields['confidence_score']} / {fields['confidence_band']}"),
        ("QA status", fields["qa_status"]),
        ("Canonical source", f"Markdown SHA-256 {source_hash[:16]}..."),
    )
    for label, value in metadata:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        _set_run_font(paragraph.add_run(f"{label}: "), size=10, bold=True)
        _set_run_font(paragraph.add_run(str(value)), size=10)

    rule = document.add_paragraph()
    rule.paragraph_format.space_before = Pt(8)
    rule.paragraph_format.space_after = Pt(10)
    p_pr = rule._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "2E74B5")
    borders.append(bottom)
    p_pr.append(borders)


def _add_table(document: DocumentObject, rows: Sequence[Sequence[str]],
               content_width_dxa: int) -> None:
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    widths = _column_widths(rows, content_width_dxa)
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            cell.text = value
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            if row_index == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "F2F4F7")
                cell._tc.get_or_add_tcPr().append(shading)
            for run in paragraph.runs:
                _set_run_font(run, size=9.5, bold=(row_index == 0))
            if column_index > 0 and re.fullmatch(r"\(?[-+]?[$£€]?\d[\d,]*(?:\.\d+)?%?\)?(?:x)?", value):
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _repeat_header(table.rows[0])
    for row in table.rows:
        _prevent_row_split(row)
    _apply_table_geometry(table, widths, content_width_dxa)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def _project_table(rows: Sequence[Sequence[str]]) -> list[list[list[str]]]:
    """Return presentation tables without dropping or changing any source cell.

    Tables through six columns stay portrait.  Seven through twelve columns use
    one landscape table.  Wider tables are split into landscape column blocks,
    repeating only the first/key column as permitted by CP_AB_EXPORT_SPEC.
    """

    column_count = len(rows[0])
    copied = [[str(value) for value in row] for row in rows]
    if column_count <= LANDSCAPE_TABLE_COLUMN_LIMIT:
        return [copied]

    projected: list[list[list[str]]] = []
    non_key_columns = list(range(1, column_count))
    chunk_size = LANDSCAPE_TABLE_COLUMN_LIMIT - 1
    for start in range(0, len(non_key_columns), chunk_size):
        selected = non_key_columns[start : start + chunk_size]
        projected.append([[row[0], *(row[index] for index in selected)] for row in copied])
    return projected


def _projected_semantic_blocks(blocks: Sequence[Block]) -> list[str]:
    serialised: list[str] = []
    for block in blocks:
        if block.kind != "table":
            serialised.extend(serialise_blocks([block]))
            continue
        for rows in _project_table(block.value):
            serialised.extend("T:" + "\t".join(row) for row in rows)
    return serialised


def _configure_section(section, *, landscape: bool) -> None:
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width = Inches(11 if landscape else 8.5)
    section.page_height = Inches(8.5 if landscape else 11)
    section.top_margin = Inches(0.75 if landscape else 1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(0.75 if landscape else 1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def _configure_running_header_footer(section, fields: dict[str, Any]) -> None:
    """Give every section explicit running matter for renderer parity."""

    section.header.is_linked_to_previous = False
    header = section.header.paragraphs[0]
    header.clear()
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    subject_id = fields.get("scope_key", fields.get("issuer_id", "UNKNOWN"))
    _set_run_font(
        header.add_run(f"{subject_id} | {fields['module_id']} | {fields['reporting_period']}"),
        size=8.5,
        color="666666",
    )

    section.footer.is_linked_to_previous = False
    footer = section.footer.paragraphs[0]
    footer.clear()
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_run_font(footer.add_run("Generated from validated canonical Markdown"), size=8, color="777777")


def build_docx(fields: dict[str, Any], blocks: Sequence[Block], output: Path,
               source_hash: str) -> None:
    """Build and deterministically serialise one Word projection."""

    document = Document()
    _configure_styles(document)
    section = document.sections[0]
    _configure_section(section, landscape=False)
    _configure_running_header_footer(section, fields)

    properties = document.core_properties
    subject_name = fields.get("subject_name", fields.get("issuer_name", "Unknown subject"))
    properties.title = f"{subject_name} {fields['module_id']}"
    properties.subject = "Deterministic projection of canonical CP Markdown handoff"
    properties.author = "CP Agents deterministic renderer"
    properties.last_modified_by = "CP Agents deterministic renderer"
    properties.created = FIXED_TIMESTAMP
    properties.modified = FIXED_TIMESTAMP

    _add_title_block(document, fields, source_hash)

    current_landscape = False
    landscape_has_content = False

    def ensure_orientation(landscape: bool) -> None:
        nonlocal current_landscape, landscape_has_content
        if current_landscape == landscape:
            return
        new_section = document.add_section(WD_SECTION.NEW_PAGE)
        _configure_section(new_section, landscape=landscape)
        _configure_running_header_footer(new_section, fields)
        current_landscape = landscape
        landscape_has_content = False

    for block in blocks:
        if block.kind == "h2":
            ensure_orientation(False)
            document.add_paragraph(block.value, style="Heading 1")
        elif block.kind == "h3":
            ensure_orientation(False)
            document.add_paragraph(block.value, style="Heading 2")
        elif block.kind == "h4":
            ensure_orientation(False)
            document.add_paragraph(block.value, style="Heading 3")
        elif block.kind == "paragraph":
            ensure_orientation(False)
            document.add_paragraph(block.value)
        elif block.kind == "bullet":
            ensure_orientation(False)
            document.add_paragraph(block.value, style="List Bullet")
        elif block.kind == "number":
            ensure_orientation(False)
            document.add_paragraph(block.value, style="List Number")
        elif block.kind == "table":
            landscape = len(block.value[0]) > WIDE_TABLE_COLUMN_LIMIT
            ensure_orientation(landscape)
            for projected_index, projected_rows in enumerate(_project_table(block.value)):
                if landscape and landscape_has_content and projected_index == 0:
                    new_section = document.add_section(WD_SECTION.NEW_PAGE)
                    _configure_section(new_section, landscape=True)
                    _configure_running_header_footer(new_section, fields)
                    current_landscape = True
                    landscape_has_content = False
                _add_table(
                    document,
                    projected_rows,
                    LANDSCAPE_CONTENT_WIDTH_DXA if landscape else PORTRAIT_CONTENT_WIDTH_DXA,
                )
                landscape_has_content = landscape
        else:
            raise ValueError(f"Unsupported block kind: {block.kind}")

    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="cp-docx-") as temporary:
        raw = Path(temporary) / "raw.docx"
        document.save(raw)
        _normalise_docx_zip(raw, output)


def _normalise_docx_zip(source: Path, output: Path) -> None:
    """Remove ZIP timestamp/order nondeterminism from python-docx output."""

    with ZipFile(source, "r") as incoming, ZipFile(output, "w", ZIP_DEFLATED, compresslevel=9) as outgoing:
        for name in sorted(incoming.namelist()):
            payload = incoming.read(name)
            info = ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            info.create_system = 3
            outgoing.writestr(info, payload)


def render_markdown(markdown: Path, output: Path) -> dict[str, Any]:
    text = markdown.read_text(encoding="utf-8")
    result = validate_text(text)
    if result.exit_code != 0:
        details = list(result.errors + result.identity_mismatches)
        if result.fields is not None and result.fields.get("qa_status") == "Blocked":
            details.append("qa_status is Blocked")
        raise ValueError("Invalid canonical Markdown: " + "; ".join(details))
    fields, body = _parse_frontmatter(text)
    blocks = parse_markdown_body(body)
    source_hash = hashlib.sha256(text.encode("utf-8")).hexdigest()
    build_docx(fields, blocks, output, source_hash)
    return {
        "source_sha256": source_hash,
        "docx_sha256": hashlib.sha256(output.read_bytes()).hexdigest(),
        "canonical_semantic_blocks": serialise_blocks(blocks),
        "semantic_blocks": _projected_semantic_blocks(blocks),
    }


def _iter_blocks(parent: DocumentObject) -> Iterable[Paragraph | Table]:
    for child in parent.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def extract_docx_body_blocks(path: Path) -> list[str]:
    """Extract the ordered Word body from Audit Summary onward."""

    document = Document(path)
    serialised: list[str] = []
    started = False
    for item in _iter_blocks(document):
        if isinstance(item, Paragraph):
            text = item.text.strip()
            if not text:
                continue
            if not started:
                if text != "Audit Summary":
                    continue
                started = True
            style = item.style.name if item.style is not None else ""
            if style == "Heading 1":
                serialised.append(f"H2:{text}")
            elif style == "Heading 2":
                serialised.append(f"H3:{text}")
            elif style == "Heading 3":
                serialised.append(f"H4:{text}")
            else:
                serialised.append(f"P:{text}")
        elif started:
            for row in item.rows:
                serialised.append("T:" + "\t".join(cell.text.strip() for cell in row.cells))
    return serialised


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("markdown", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args(argv)
    metadata = render_markdown(args.markdown, args.output)
    extracted = extract_docx_body_blocks(args.output)
    if extracted != metadata["semantic_blocks"]:
        print("PARITY_FAILURE")
        return 2
    print(f"VALID source_sha256={metadata['source_sha256']} docx_sha256={metadata['docx_sha256']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
